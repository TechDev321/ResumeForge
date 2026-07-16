from __future__ import annotations

import argparse
import os
import sys
from pathlib import Path

from dotenv import load_dotenv
from docx import Document

from resume_generator.cover_letter_template import create_cover_letter_template
from resume_generator.docx_utils import iter_paragraphs
from resume_generator.openai_cover_letter import generate_cover_letter_replacements
from resume_generator.template_engine import apply_replacements, extract_placeholders


def _read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def _read_docx_text(path: Path) -> str:
    doc = Document(str(path))
    return "\n".join(p.text for p in iter_paragraphs(doc) if (p.text or "").strip())


def _find_latest_resume(out_dir: Path) -> Path | None:
    candidates = [
        p
        for p in out_dir.glob("*.docx")
        if "Cover_Letter" not in p.name and "cover_letter" not in p.name.lower()
    ]
    if not candidates:
        return None
    return max(candidates, key=lambda p: p.stat().st_mtime)


def _ensure_template(template_path: Path) -> None:
    if not template_path.exists():
        create_cover_letter_template(template_path)


def main() -> int:
    load_dotenv()

    ap = argparse.ArgumentParser(description="Generate a cover letter docx from a resume + job description")
    ap.add_argument("--template", default="corey_cover_letter_template.docx", help="Path to cover letter template .docx")
    ap.add_argument("--prompt", default="Cover Letter Prompt.txt", help="Path to cover letter prompt/rules text")
    ap.add_argument("--jd", default="jd_example.txt", help="Path to job description text file")
    ap.add_argument(
        "--resume",
        default=None,
        help="Path to generated resume .docx (default: newest .docx in --out, excluding cover letters)",
    )
    ap.add_argument("--out", default=".", help="Output directory")
    ap.add_argument("--model", default=os.getenv("OPENAI_MODEL", "gpt-4.1-mini"))
    ap.add_argument("--no-cache", action="store_true", help="Disable caching")

    args = ap.parse_args()

    api_key = os.getenv("OPENAI_API_KEY")
    if not api_key:
        print("ERROR: OPENAI_API_KEY is not set. Put it in your environment or a .env file.", file=sys.stderr)
        return 2

    template_path = Path(args.template)
    prompt_path = Path(args.prompt)
    jd_path = Path(args.jd)
    out_dir = Path(args.out)
    out_dir.mkdir(parents=True, exist_ok=True)

    _ensure_template(template_path)

    resume_path = Path(args.resume) if args.resume else _find_latest_resume(out_dir)
    if resume_path is None or not resume_path.exists():
        print(
            "ERROR: No resume .docx found. Generate one first with generate_resume.py, "
            "then pass --resume PATH or place the resume in --out.",
            file=sys.stderr,
        )
        return 2

    prompt_rules = _read_text(prompt_path)
    jd_text = _read_text(jd_path)
    resume_text = _read_docx_text(resume_path)

    doc = Document(str(template_path))
    placeholders = extract_placeholders(doc)
    if not placeholders:
        print("ERROR: No placeholders found in cover letter template.", file=sys.stderr)
        return 2

    cache_dir = None
    cache_enabled = (os.getenv("COVER_LETTER_CACHE", "1") == "1") and (not args.no_cache)
    if cache_enabled:
        cache_dir = Path("cache") / "cover_letter"

    file_name, replacements = generate_cover_letter_replacements(
        api_key=api_key,
        model=args.model,
        prompt_rules=prompt_rules,
        jd_text=jd_text,
        resume_text=resume_text,
        placeholders=placeholders,
        cache_dir=cache_dir,
    )

    apply_replacements(doc, replacements)
    out_path = out_dir / file_name
    doc.save(str(out_path))
    print(str(out_path))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
